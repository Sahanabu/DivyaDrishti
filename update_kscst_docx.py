import os
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC_PATH = r"C:\Users\ratan\Desktop\DivyaDrishti\03_KSCST_SPP_50S_Proposal_Format_2026-27.docx"
OUT_PATH = r"C:\Users\ratan\Desktop\DivyaDrishti\03_KSCST_SPP_50S_Proposal_Format_2026-27_Updated.docx"

doc = docx.Document(SRC_PATH)

def format_cell_paragraphs(cell, text_entries):
    """
    Replaces all paragraphs in a cell with the given list of text entries.
    Preserves any nested tables inside the cell.
    """
    # Clear existing paragraphs
    p0 = cell.paragraphs[0]
    p0.text = ""
    # Remove remaining paragraphs
    for p in list(cell.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    
    for i, line in enumerate(text_entries):
        p = cell.add_paragraph() if i > 0 else p0
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        is_bold_header = any(line.startswith(h) for h in [
            "Background and Need", "Review of Existing Work", "Problem Statement", 
            "Objectives", "Innovation / Novelty", "1. WHAT IS NEW", "2. HOW IS IT DIFFERENT", "3. WHY IS IT INNOVATIVE",
            "1. EXISTING TECHNOLOGIES", "2. COMMERCIAL PRODUCTS", "3. RESEARCH PAPERS", "4. PATENTS", "5. EXISTING SOLUTIONS",
            "TECHNICAL", "SOCIAL", "ENVIRONMENTAL", "ECONOMIC", "INDUSTRIAL", "ACADEMIC", "POTENTIAL USERS",
            "Expected Deliverables", "Expected Outcomes", "Relevance", "Major Requirements", "References",
            "1. HARDWARE", "2. SOFTWARE", "3. CONSUMABLES", "4. FIELD STUDIES"
        ])
        
        r = p.add_run(line)
        r.font.name = "Arial" if is_bold_header else "Times New Roman"
        r.font.size = Pt(10)
        if is_bold_header:
            r.bold = True

print("Script template ready.")
